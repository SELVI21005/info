from datetime import datetime
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def generate_report(analysis, alert=None):
    return {
        "report_type": "Malware Analysis Report",
        "generated_at": datetime.now().isoformat(),
        "filename": analysis.get("filename"),
        "classification": analysis.get("classification"),
        "risk_score": analysis.get("risk_score"),
        "risk_level": analysis.get("risk_level"),
        "risk_explanation": analysis.get("risk_explanation", []),
        "sha256": analysis.get("sha256"),
        "file_size": analysis.get("file_size"),
        "static_analysis": analysis.get("static_analysis"),
        "ml_analysis": analysis.get("ml_analysis"),
        "alert_status": {
            "generated": bool(alert),
            "reason": alert.get("message") if alert else "No alert generated",
            "timestamp": alert.get("timestamp") if alert else None,
        },
    }


def _display_value(value, max_length=900):
    if value is None or value == "":
        return "N/A"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, indent=2, default=str)
    else:
        value = str(value)

    if len(value) > max_length:
        return value[:max_length] + "\n[truncated for report layout]"

    return value


def _section_rows(report):
    static = report.get("static_analysis") or {}
    identification = static.get("file_identification") or {}
    metadata = static.get("metadata") or {}
    pe = static.get("pe_analysis") or {}
    ml = report.get("ml_analysis") or {}
    alert = report.get("alert_status") or {}

    return [
        ("FILE INFORMATION", [
            ("File name", report.get("filename")),
            ("File type", identification.get("file_type")),
            ("File size", report.get("file_size")),
            ("File identification", identification),
            ("MD5", static.get("md5")),
            ("SHA-256", report.get("sha256") or static.get("sha256")),
            ("Metadata", metadata),
        ]),
        ("CLASSIFICATION SUMMARY", [
            ("Final classification", report.get("classification")),
            ("ML classification", ml.get("classification")),
            ("ML confidence", ml.get("confidence")),
            ("Static classification", static.get("static_classification")),
            ("Risk score", report.get("risk_score")),
            ("Risk level", report.get("risk_level")),
        ]),
        ("STATIC ANALYSIS", [
            ("Strings", static.get("strings_count")),
            ("Suspicious keywords", static.get("suspicious_keywords")),
            ("URLs", static.get("urls")),
            ("IP addresses", static.get("ips")),
            ("PE information", pe),
            ("Sections", pe.get("number_of_sections")),
            ("Imports/APIs", pe.get("imports") or static.get("suspicious_imports")),
            ("YARA analysis", static.get("yara_analysis")),
            ("Static risk score", static.get("static_risk_score")),
        ]),
        ("MACHINE LEARNING ANALYSIS", [
            ("Model status", ml.get("status")),
            ("Model name", ml.get("model_name")),
            ("Feature count", ml.get("feature_count")),
            ("Classification", ml.get("classification")),
            ("Confidence", ml.get("confidence")),
        ]),
        ("RISK ASSESSMENT", [
            ("Final risk score", report.get("risk_score")),
            ("Risk level", report.get("risk_level")),
            ("Risk explanation", report.get("risk_explanation")),
        ]),
        ("ALERT STATUS", [
            ("Alert generated", alert.get("generated", False)),
            ("Alert reason", alert.get("reason")),
            ("Alert timestamp", alert.get("timestamp")),
        ]),
    ]


def build_pdf_report(report):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#12304a"),
        spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="ReportSection",
        parent=styles["Heading2"],
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#087f9b"),
        spaceBefore=13,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="ReportBody",
        parent=styles["BodyText"],
        fontSize=8,
        leading=10,
    ))

    story = [
        Paragraph("ThreatLens AI", styles["ReportTitle"]),
        Paragraph("Malware Analysis &amp; Threat Detection Report", styles["Heading2"]),
        Spacer(1, 8),
        Table([
            ["Report ID", _display_value(report.get("id"))],
            ["Generated", _display_value(report.get("generated_at"))],
        ], colWidths=[1.35 * inch, 5.9 * inch], style=TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e8f4f7")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8d9e0")),
            ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d6e6ea")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("PADDING", (0, 0), (-1, -1), 6),
        ])),
    ]

    for heading, rows in _section_rows(report):
        table_rows = [[Paragraph("Field", styles["ReportBody"]), Paragraph("Value", styles["ReportBody"])]]
        for label, value in rows:
            formatted = _display_value(value).replace("\n", "<br/>")
            table_rows.append([
                Paragraph(str(label), styles["ReportBody"]),
                Paragraph(formatted.replace("&", "&amp;"), styles["ReportBody"]),
            ])
        story.append(Paragraph(heading, styles["ReportSection"]))
        story.append(Table(table_rows, colWidths=[1.65 * inch, 5.6 * inch], repeatRows=1, style=TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#12304a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5df")),
            ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8fbfc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("PADDING", (0, 0), (-1, -1), 5),
        ])))

    document.build(story)
    output.seek(0)
    return output


def build_excel_report(report):
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    header_fill = PatternFill("solid", fgColor="12304A")
    section_fill = PatternFill("solid", fgColor="DDEFF3")

    for sheet_name, rows in [
        ("Summary", [("Report ID", report.get("id")), ("Generated", report.get("generated_at")), ("Classification", report.get("classification")), ("Risk Score", report.get("risk_score")), ("Risk Level", report.get("risk_level"))]),
        ("File Information", _section_rows(report)[0][1]),
        ("Static Analysis", _section_rows(report)[2][1]),
        ("ML Analysis", _section_rows(report)[3][1]),
        ("Risk Assessment", _section_rows(report)[4][1]),
        ("Indicators", [("Suspicious keywords", (report.get("static_analysis") or {}).get("suspicious_keywords")), ("URLs", (report.get("static_analysis") or {}).get("urls")), ("IP addresses", (report.get("static_analysis") or {}).get("ips")), ("Suspicious imports", (report.get("static_analysis") or {}).get("suspicious_imports"))]),
    ]:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["Field", "Value"])
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
        for label, value in rows:
            sheet.append([label, _display_value(value)])
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 26
        sheet.column_dimensions["B"].width = 90
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row_number in range(2, sheet.max_row + 1):
            if row_number % 2 == 0:
                for cell in sheet[row_number]:
                    cell.fill = section_fill

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output